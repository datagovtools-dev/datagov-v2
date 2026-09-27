"""Re-apply updated _assess_standard_format logic to all projects.

Reads persisted source files from /app/uploads, re-parses ALL rows
(not just the 100-row sample cap used at import time), then updates
standard_format and distinct_values for every matching metadata record.

Usage (inside container):
    python scripts/reapply_standard_format.py
    python scripts/reapply_standard_format.py --dry-run
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from typing import Any

import openpyxl
from sqlalchemy import select

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from app.database import get_sync_session  # noqa: E402
from app.models.metadata import MetadataRecord, ProjectSourceFile  # noqa: E402

# ── Inline copies of the updated helpers (no import path issues) ──────────────

_CAT_NAME_HINTS = {
    "type", "status", "level", "category", "cat", "grade", "tier", "segment",
    "class", "flag", "brand", "color", "colour", "region", "zone", "dept",
    "channel", "method", "mode", "rank", "priority", "state", "group", "grp",
}
_TEXT_NAME_HINTS = {
    "name", "description", "desc", "notes", "note", "remark", "comment",
    "address", "addr", "text", "message", "msg", "content", "detail", "info",
}


def _assess_standard_format(values: list[Any], column_name: str | None = None) -> str | None:
    non_null = [
        str(v).strip() for v in values
        if v is not None and str(v).strip() and str(v).strip().lower() not in ("nan", "none", "")
    ]
    if not non_null:
        return None

    total = len(non_null)
    unique_vals = list(dict.fromkeys(non_null))
    n_unique = len(unique_vals)

    bool_set = {"true", "false", "yes", "no", "y", "n", "0", "1", "t", "f"}
    bool_positive = {"true", "yes", "y", "1", "t"}
    bool_unique_lower = {v.lower() for v in unique_vals}
    if n_unique <= 4 and bool_unique_lower <= bool_set:
        seen_lower: set = set()
        deduped: list = []
        for v in unique_vals:
            if v.lower() not in seen_lower:
                seen_lower.add(v.lower())
                deduped.append(v)
        ordered = sorted(deduped, key=lambda v: (0 if v.lower() in bool_positive else 1))
        return f"Boolean ({' / '.join(ordered)})"

    all_numeric = all(re.match(r"^-?\d+(\.\d+)?$", v) for v in unique_vals)
    if not all_numeric:
        col_parts = set(re.split(r"[_\s\-]", (column_name or "").lower()))
        cat_hint  = bool(col_parts & _CAT_NAME_HINTS)
        text_hint = bool(col_parts & _TEXT_NAME_HINTS)

        avg_val_len = sum(len(v) for v in non_null) / total
        long_values = avg_val_len > 35

        n_thresh = 20 if cat_hint else 15
        r_thresh = 0.12 if cat_hint else 0.10

        avg_freq = total / n_unique
        if (not long_values and not text_hint and
            (n_unique <= n_thresh
             or (total >= 20 and n_unique / total < r_thresh)
             or (total >= 50 and avg_freq >= 3.0 and n_unique <= 30 and avg_val_len <= 30))):
            cats = sorted(unique_vals[:10])
            suffix = ", ..." if n_unique > 10 else ""
            return f"Category: {', '.join(cats)}{suffix}"

    date_patterns = [
        (r"^\d{4}-\d{2}-\d{2}$",   "Date (YYYY-MM-DD)"),
        (r"^\d{2}/\d{2}/\d{4}$",   "Date (DD/MM/YYYY)"),
        (r"^\d{2}-\d{2}-\d{4}$",   "Date (DD-MM-YYYY)"),
        (r"^\d{4}/\d{2}/\d{2}$",   "Date (YYYY/MM/DD)"),
    ]
    for pattern, label in date_patterns:
        if sum(1 for v in non_null if re.match(pattern, v)) / total >= 0.85:
            return label

    if sum(1 for v in non_null if re.match(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}", v)) / total >= 0.85:
        return "Datetime (YYYY-MM-DD HH:MM:SS)"

    if sum(1 for v in non_null if re.match(r"^[\w.+\-]+@[\w\-]+\.[a-zA-Z]{2,}$", v)) / total >= 0.8:
        return "Email (name@domain.com)"

    if sum(
        1 for v in non_null
        if re.match(r"^[+\d][\d\s\-().x]{6,20}$", v)
        and re.search(r"[+\s\-().x]", v)
    ) / total >= 0.8:
        return "Phone number"

    if sum(1 for v in non_null if re.match(r"^-?\d+$", v)) / total >= 0.9:
        return "Integer (whole number)"

    decimal_matches = [v for v in non_null if re.match(r"^-?\d+\.\d+$", v)]
    if len(decimal_matches) / total >= 0.8:
        dp_counts = [len(v.split(".")[1]) for v in decimal_matches]
        if len(set(dp_counts)) == 1:
            return f"Decimal ({dp_counts[0]} decimal places)"
        return "Decimal number"

    id_matches = sum(
        1 for v in non_null
        if re.match(r"^[A-Z]{2,}[-_]\d+$", v)
        or re.match(r"^[A-Z]{2,}\d{2,}$", v)
        or re.match(r"^\d{6,20}$", v)
    )
    if id_matches / total >= 0.8:
        return f"ID / Code (e.g. {non_null[0]})"

    avg_len = sum(len(v) for v in non_null) / total
    if avg_len > 40:
        return "Free text (long description)"
    return "Free text"


def _get_distinct_values(values: list[Any], standard_format: str | None) -> str | None:
    non_null = sorted({str(v).strip() for v in values if v is not None and str(v).strip()})
    if not non_null:
        return None
    is_cat_bool = standard_format and (
        standard_format.startswith("Category:") or standard_format.startswith("Boolean")
    )
    if not is_cat_bool and len(non_null) > 25:
        return None
    return ", ".join(non_null[:20]) if non_null else None


# ── File parsing ──────────────────────────────────────────────────────────────

def parse_excel_all_rows(file_path: str) -> dict[str, dict[str, list[Any]]]:
    """Return {sheet_name: {column: [all values]}} — no row cap."""
    result: dict[str, dict[str, list[Any]]] = {}
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                continue
            headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
            data: dict[str, list[Any]] = {h: [] for h in headers}
            for row in rows[1:]:
                for h, v in zip(headers, row):
                    data[h].append(v)
            result[sheet_name] = data
    finally:
        wb.close()
    return result


# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    db = get_sync_session()

    # Latest file per (project_id, original_filename)
    latest: dict[tuple, ProjectSourceFile] = {}
    for f in db.scalars(select(ProjectSourceFile).order_by(ProjectSourceFile.uploaded_at.desc())):
        latest.setdefault((f.project_id, f.original_filename), f)
    files = list(latest.values())

    if not files:
        print("No persisted source files found — upload via the Metadata menu first.")
        sys.exit(0)

    print(f"{'DRY RUN — ' if args.dry_run else ''}Processing {len(files)} source files …\n")

    total_updated = total_skipped = total_no_change = 0

    for f in files:
        proj_id = f.project_id
        orig_name = f.original_filename
        path = f.stored_path

        if not os.path.exists(path):
            print(f"  [MISSING] {path}")
            continue

        print(f"  Parsing  {orig_name} …")
        try:
            sheets = parse_excel_all_rows(path)
        except Exception as exc:
            print(f"  [ERROR]   {orig_name}: {exc}")
            continue

        for sheet_name, columns in sheets.items():
            domain_table = f"{orig_name} - {sheet_name}"

            for col, values in columns.items():
                new_fmt = _assess_standard_format(values, col)
                _PHONE_KEYWORDS = {"phone", "mobile", "tel", "hp", "handphone", "telepon", "nohp", "no_hp"}
                if new_fmt in ("Integer (whole number)", "Free text", None):
                    if any(kw in col.lower() for kw in _PHONE_KEYWORDS):
                        new_fmt = "Phone number"
                new_dv  = _get_distinct_values(values, new_fmt)

                row = db.scalars(select(MetadataRecord).where(
                    MetadataRecord.project_id == proj_id,
                    MetadataRecord.data_domain_table == domain_table,
                    MetadataRecord.data_attribute == col,
                )).first()

                if not row:
                    total_skipped += 1
                    continue

                old_fmt = row.standard_format
                old_dv  = row.distinct_values

                changed = (new_fmt != old_fmt) or (new_dv != old_dv)
                tag = f"{old_fmt!r}  →  {new_fmt!r}" if changed else "(no change)"
                print(f"    [{domain_table}] {col:40s}  {tag}")

                if changed:
                    if not args.dry_run:
                        row.standard_format = new_fmt
                        row.distinct_values = new_dv
                    total_updated += 1
                else:
                    total_no_change += 1

    if not args.dry_run:
        db.commit()

    print(f"\n{'DRY RUN — ' if args.dry_run else ''}Done.")
    print(f"  Updated   : {total_updated}")
    print(f"  No change : {total_no_change}")
    print(f"  Skipped   : {total_skipped} (column not found in metadata_records)")

    db.close()


if __name__ == "__main__":
    main()
