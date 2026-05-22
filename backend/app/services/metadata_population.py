"""Metadata auto-population helpers shared by API routes and workers."""
from __future__ import annotations

import csv
import os
import re
import tempfile
from datetime import date, datetime, timezone
from typing import Any, Sequence
from uuid import UUID

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.metadata import MetadataRecord

SAMPLE_ROW_LIMIT = 100

_PII_KEYWORDS = {
    "nama", "telepon", "nik", "ktp", "passport", "gaji", "agama", "kesehatan",
    "alamat", "tanggal_lahir", "tgl_lahir", "jenis_kelamin", "no_rekening",
    "npwp", "sim", "foto", "wajah", "sidik_jari",
    "name", "phone", "mobile", "email", "mail", "salary", "gender", "religion",
    "health", "address", "birthdate", "birth_date", "dob", "ssn", "passport",
    "account_number", "credit_card", "card_number", "face", "fingerprint",
    "biometric", "race", "ethnicity", "nationality", "income", "tax_id",
}

_ABBREV_MAP = {
    "id": "Identifier", "nm": "Name", "cd": "Code", "dt": "Date",
    "dttm": "Date Time", "no": "Number", "num": "Number", "qty": "Quantity",
    "amt": "Amount", "val": "Value", "desc": "Description", "flg": "Flag",
    "ind": "Indicator", "sts": "Status", "typ": "Type", "cat": "Category",
    "grp": "Group", "seq": "Sequence", "ref": "Reference", "src": "Source",
    "tgt": "Target", "tot": "Total", "avg": "Average", "max": "Maximum",
    "min": "Minimum", "cnt": "Count", "pct": "Percentage", "ts": "Timestamp",
    "upd": "Updated", "cre": "Created", "del": "Deleted", "usr": "User",
    "org": "Organisation", "dept": "Department", "div": "Division",
    "loc": "Location", "addr": "Address", "tel": "Telephone", "fax": "Fax",
    "url": "URL", "img": "Image", "doc": "Document", "file": "File",
    "txt": "Text", "msg": "Message", "err": "Error", "log": "Log",
}


def _classify_sensitivity(attribute: str) -> str:
    lower = attribute.lower()
    parts = re.split(r"[_\s\-]", lower)
    for part in parts:
        if part in _PII_KEYWORDS:
            return "Highly Confidential"
    for keyword in _PII_KEYWORDS:
        if keyword in lower:
            return "Highly Confidential"
    return "Confidential"


def _expand_business_term(attribute: str) -> str:
    clean = re.sub(r"^(tbl_|fk_|idx_|pk_|f_|t_)", "", attribute, flags=re.IGNORECASE)
    parts = re.split(r"[_\s\-]", clean)
    expanded = []
    for part in parts:
        lower = part.lower()
        expanded.append(_ABBREV_MAP.get(lower, part.capitalize()))
    return " ".join(expanded)


def _detect_data_type(values: list[Any]) -> str:
    samples = [v for v in values if v is not None][:20]
    if not samples:
        return "STRING"
    type_map = {
        "BOOLEAN": lambda v: str(v).lower() in ("true", "false", "0", "1"),
        "INTEGER": lambda v: str(v).lstrip("-").isdigit(),
        "FLOAT": lambda v: re.match(r"^-?\d+\.\d+$", str(v)) is not None,
        "DATE": lambda v: re.match(r"^\d{4}-\d{2}-\d{2}$", str(v)) is not None,
        "DATETIME": lambda v: re.match(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}", str(v)) is not None,
        "TIME": lambda v: re.match(r"^\d{2}:\d{2}(:\d{2})?$", str(v)) is not None,
    }
    for dtype, check in type_map.items():
        if sum(1 for sample in samples if check(sample)) / len(samples) >= 0.8:
            return dtype
    return "STRING"


def _get_sample_data(values: list[Any]) -> str:
    for value in values:
        if value is not None and str(value).strip():
            return str(value)[:200]
    return "(All Blank)"


def _is_primary_key_candidate(values: list[Any]) -> bool:
    non_null = [value for value in values if value is not None]
    if not non_null:
        return False
    return len(set(str(value) for value in non_null)) == len(non_null)


def _is_nullable(values: list[Any]) -> bool:
    return any(value is None or str(value).strip() == "" for value in values)


def _get_distinct_values(values: list[Any], standard_format: str | None) -> str | None:
    """Return comma-separated distinct values for categorical/boolean columns only."""
    if not standard_format:
        return None
    if not (standard_format.startswith("Category:") or standard_format.startswith("Boolean")):
        return None
    non_null = sorted({str(v).strip() for v in values if v is not None and str(v).strip()})
    return ", ".join(non_null[:20]) if non_null else None


def _assess_standard_format(values: list[Any]) -> str | None:
    non_null = [
        str(value).strip() for value in values
        if value is not None and str(value).strip() and str(value).strip().lower() not in ("nan", "none", "")
    ]
    if not non_null:
        return None

    total = len(non_null)
    unique_values = list(dict.fromkeys(non_null))
    unique_count = len(unique_values)

    bool_set = {"true", "false", "yes", "no", "y", "n", "0", "1", "t", "f"}
    if unique_count <= 4 and all(value.lower() in bool_set for value in unique_values):
        return "Boolean (Yes/No or True/False)"

    if unique_count <= 15 or (total >= 20 and unique_count / total < 0.1):
        categories = sorted(unique_values[:10])
        suffix = ", ..." if unique_count > 10 else ""
        return f"Category: {', '.join(categories)}{suffix}"

    date_patterns = [
        (r"^\d{4}-\d{2}-\d{2}$", "Date (YYYY-MM-DD)"),
        (r"^\d{2}/\d{2}/\d{4}$", "Date (DD/MM/YYYY)"),
        (r"^\d{2}-\d{2}-\d{4}$", "Date (DD-MM-YYYY)"),
        (r"^\d{4}/\d{2}/\d{2}$", "Date (YYYY/MM/DD)"),
    ]
    for pattern, label in date_patterns:
        if sum(1 for value in non_null if re.match(pattern, value)) / total >= 0.85:
            return label

    if sum(1 for value in non_null if re.match(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}", value)) / total >= 0.85:
        return "Datetime (YYYY-MM-DD HH:MM:SS)"

    if sum(1 for value in non_null if re.match(r"^[\w.+\-]+@[\w\-]+\.[a-zA-Z]{2,}$", value)) / total >= 0.8:
        return "Email (name@domain.com)"

    if sum(1 for value in non_null if re.match(r"^[+\d][\d\s\-().]{6,18}$", value)) / total >= 0.8:
        return "Phone number"

    if sum(1 for value in non_null if re.match(r"^-?\d+$", value)) / total >= 0.9:
        return "Integer (whole number)"

    decimal_matches = [value for value in non_null if re.match(r"^-?\d+\.\d+$", value)]
    if len(decimal_matches) / total >= 0.8:
        decimal_places = [len(value.split(".")[1]) for value in decimal_matches]
        if len(set(decimal_places)) == 1:
            return f"Decimal ({decimal_places[0]} decimal places)"
        return "Decimal number"

    id_matches = sum(
        1 for value in non_null
        if re.match(r"^[A-Z]{2,}[-_]\d+$", value) or re.match(r"^\d{6,20}$", value)
    )
    if id_matches / total >= 0.8:
        return f"ID / Code (e.g. {non_null[0]})"

    average_length = sum(len(value) for value in non_null) / total
    if average_length > 40:
        return "Free text (long description)"
    return "Free text"


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value


def _make_unique_headers(raw_headers: Sequence[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, raw in enumerate(raw_headers):
        header = str(raw).strip() if raw is not None and str(raw).strip() else f"col_{index + 1}"
        count = seen.get(header, 0) + 1
        seen[header] = count
        headers.append(header if count == 1 else f"{header}_{count}")
    return headers


def _normalise_row(raw_row: Sequence[Any], width: int) -> list[Any]:
    row = list(raw_row)
    return [_clean_cell(row[index]) if index < len(row) else None for index in range(width)]


def parse_csv_upload_content(content: bytes, filename: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Parse CSV bytes into sheet metadata and a bounded sample payload."""
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1", errors="replace")

    rows = list(csv.reader(text.splitlines()))
    sheet_name = os.path.splitext(filename or "data")[0]
    if not rows:
        return ([{"sheet_name": sheet_name, "column_count": 0, "row_count": 0}], [])

    headers = _make_unique_headers(rows[0])
    sample_rows: list[list[Any]] = []
    row_count = 0
    for raw_row in rows[1:]:
        row_count += 1
        if len(sample_rows) < SAMPLE_ROW_LIMIT:
            sample_rows.append(_normalise_row(raw_row, len(headers)))

    sheets = [{"sheet_name": sheet_name, "column_count": len(headers), "row_count": row_count}]
    tables = [{
        "table_name": sheet_name,
        "columns": headers,
        "sample_rows": sample_rows,
        "row_count": row_count,
    }]
    return sheets, tables


def parse_excel_upload_file(file_path: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Parse an Excel workbook into sheet metadata and bounded sample payloads."""
    import openpyxl

    sheets: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows_iter = worksheet.iter_rows(values_only=True)
            header_row = next(rows_iter, None)
            if not header_row:
                sheets.append({"sheet_name": sheet_name, "column_count": 0, "row_count": 0})
                continue

            headers = _make_unique_headers(header_row)
            sample_rows: list[list[Any]] = []
            row_count = 0
            for raw_row in rows_iter:
                row_count += 1
                if len(sample_rows) < SAMPLE_ROW_LIMIT:
                    sample_rows.append(_normalise_row(raw_row, len(headers)))

            sheets.append({"sheet_name": sheet_name, "column_count": len(headers), "row_count": row_count})
            tables.append({
                "table_name": sheet_name,
                "columns": headers,
                "sample_rows": sample_rows,
                "row_count": row_count,
            })
    finally:
        workbook.close()
    return sheets, tables


def _table_value(table: Any, key: str, default: Any = None) -> Any:
    if isinstance(table, dict):
        return table.get(key, default)
    return getattr(table, key, default)


def build_tables_from_uploaded_payload(
    uploaded_tables: Sequence[Any] | None,
    table_names: Sequence[str] | None = None,
) -> tuple[dict[str, dict[str, list[Any]]], dict[str, int]]:
    """Convert uploaded table payloads into column-oriented samples."""
    selected = {name for name in (table_names or []) if name}
    tables_data: dict[str, dict[str, list[Any]]] = {}
    row_counts: dict[str, int] = {}

    for table in uploaded_tables or []:
        table_name = str(_table_value(table, "table_name", "") or "").strip()
        if not table_name or (selected and table_name not in selected):
            continue

        columns = [str(column).strip() for column in (_table_value(table, "columns", []) or []) if str(column).strip()]
        if not columns:
            continue

        data = {column: [] for column in columns}
        sample_rows = _table_value(table, "sample_rows", []) or []
        for sample_row in sample_rows:
            if isinstance(sample_row, dict):
                values = [sample_row.get(column) for column in columns]
            else:
                values = list(sample_row)
            for index, column in enumerate(columns):
                data[column].append(values[index] if index < len(values) else None)

        tables_data[table_name] = data
        try:
            row_counts[table_name] = int(_table_value(table, "row_count", 0) or 0)
        except (TypeError, ValueError):
            row_counts[table_name] = max((len(values) for values in data.values()), default=0)

    return tables_data, row_counts


def _load_excel_tables_from_temp(
    temp_file_key: str | None,
    temp_file_keys: Sequence[str] | None,
    file_names: Sequence[str] | None,
    table_names: Sequence[str] | None,
) -> tuple[dict[str, dict[str, list[Any]]], dict[str, int]]:
    keys = list(temp_file_keys or ([temp_file_key] if temp_file_key else []))
    if not keys:
        return {}, {}

    payload_tables: list[dict[str, Any]] = []
    multiple_files = len(keys) > 1
    for index, key in enumerate(keys):
        file_path = os.path.join(tempfile.gettempdir(), key)
        if not os.path.exists(file_path):
            raise FileNotFoundError("Uploaded file cache expired. Upload the file again and click Proceed Metadata.")

        original_name = file_names[index] if file_names and index < len(file_names) else key
        extension = os.path.splitext(key)[1].lower()
        if extension == ".csv":
            with open(file_path, "rb") as handle:
                _, tables = parse_csv_upload_content(handle.read(), original_name)
        else:
            _, tables = parse_excel_upload_file(file_path)

        for table in tables:
            sheet_name = table["table_name"]
            table["table_name"] = f"{original_name} - {sheet_name}" if multiple_files else sheet_name
            payload_tables.append(table)

    return build_tables_from_uploaded_payload(payload_tables, table_names)


def _load_postgresql_tables(
    connection_string: str,
    pg_schema: str,
    table_names: Sequence[str] | None,
) -> tuple[dict[str, dict[str, list[Any]]], dict[str, int]]:
    import psycopg2
    from psycopg2 import sql

    tables_data: dict[str, dict[str, list[Any]]] = {}
    row_counts: dict[str, int] = {}
    source_connection = psycopg2.connect(connection_string)
    try:
        source_cursor = source_connection.cursor()
        target_tables = list(table_names or [])
        if not target_tables:
            source_cursor.execute(
                "SELECT table_name FROM information_schema.tables "
                "WHERE table_schema = %s AND table_type = 'BASE TABLE' ORDER BY table_name",
                (pg_schema,),
            )
            target_tables = [row[0] for row in source_cursor.fetchall()]

        for requested_table in target_tables:
            table_name = requested_table.split(".", 1)[1] if requested_table.startswith(f"{pg_schema}.") else requested_table
            source_cursor.execute(
                "SELECT column_name FROM information_schema.columns "
                "WHERE table_schema = %s AND table_name = %s ORDER BY ordinal_position",
                (pg_schema, table_name),
            )
            columns = [row[0] for row in source_cursor.fetchall()]
            if not columns:
                continue

            try:
                source_cursor.execute(
                    sql.SQL("SELECT COUNT(*) FROM {}.{}").format(
                        sql.Identifier(pg_schema), sql.Identifier(table_name)
                    )
                )
                row_counts[table_name] = int(source_cursor.fetchone()[0])
            except Exception:
                row_counts[table_name] = 0

            try:
                source_cursor.execute(
                    sql.SQL("SELECT * FROM {}.{} LIMIT 100").format(
                        sql.Identifier(pg_schema), sql.Identifier(table_name)
                    )
                )
                rows = source_cursor.fetchall()
            except Exception:
                rows = []

            data = {column: [] for column in columns}
            for row in rows:
                for index, column in enumerate(columns):
                    data[column].append(row[index] if index < len(row) else None)
            tables_data[table_name] = data
    finally:
        source_connection.close()

    return tables_data, row_counts


def _load_gcp_tables(
    gcp_project: str,
    bq_dataset: str,
    table_names: Sequence[str] | None,
) -> tuple[dict[str, dict[str, list[Any]]], dict[str, int]]:
    from google.cloud import bigquery

    client = bigquery.Client(project=gcp_project)
    target_tables = list(table_names or [table.table_id for table in client.list_tables(f"{gcp_project}.{bq_dataset}")])
    tables_data: dict[str, dict[str, list[Any]]] = {}
    row_counts: dict[str, int] = {}
    for table_name in target_tables:
        ref = f"{gcp_project}.{bq_dataset}.{table_name}"
        table = client.get_table(ref)
        schema_columns = [field.name for field in table.schema]
        data = {column: [] for column in schema_columns}
        for row in client.list_rows(ref, max_results=SAMPLE_ROW_LIMIT):
            for column in schema_columns:
                data[column].append(row[column])
        if data:
            tables_data[table_name] = data
            row_counts[table_name] = int(table.num_rows or 0)
    return tables_data, row_counts


async def populate_metadata_records(
    db: AsyncSession,
    *,
    project_id: UUID,
    source_type: str,
    project_name: str,
    project_year: int,
    customer_name: str,
    line_of_business: str | None,
    owner_info: dict[str, str] | None,
    initiated_by: str,
    table_names: Sequence[str] | None = None,
    uploaded_tables: Sequence[Any] | None = None,
    temp_file_key: str | None = None,
    temp_file_keys: Sequence[str] | None = None,
    file_names: Sequence[str] | None = None,
    gcp_project: str | None = None,
    bq_dataset: str | None = None,
    connection_string: str | None = None,
    pg_schema: str = "public",
) -> dict[str, Any]:
    """Synchronously populate MetadataRecord rows for selected source tables."""
    selected_tables = list(table_names or [])
    source_type = source_type.lower()

    if source_type == "excel":
        tables_data, row_counts = build_tables_from_uploaded_payload(uploaded_tables, selected_tables)
        if not tables_data:
            tables_data, row_counts = _load_excel_tables_from_temp(
                temp_file_key, temp_file_keys, file_names, selected_tables
            )
    elif source_type == "postgresql" and connection_string:
        tables_data, row_counts = _load_postgresql_tables(connection_string, pg_schema, selected_tables)
    elif source_type == "gcp" and gcp_project and bq_dataset:
        tables_data, row_counts = _load_gcp_tables(gcp_project, bq_dataset, selected_tables)
    else:
        tables_data, row_counts = {}, {}

    if not tables_data:
        raise ValueError("No source columns were found. Upload or discover a source table, select it, then proceed again.")

    max_seq_result = await db.execute(
        select(func.coalesce(func.max(MetadataRecord.seq_no), 0)).where(MetadataRecord.project_id == project_id)
    )
    next_seq = int(max_seq_result.scalar_one() or 0)

    created = 0
    updated = 0
    owner_info = owner_info or {}
    owner_str = owner_info.get("data_owner", "")
    steward_str = (
        owner_info.get("data_steward")
        or owner_info.get("lead_business_steward")
        or owner_info.get("business_steward")
        or owner_info.get("lead_it_steward")
        or owner_info.get("it_steward")
        or ""
    )
    current_year = datetime.now(timezone.utc).year
    today = date.today()

    for table_name, columns in tables_data.items():
        if source_type == "gcp" and bq_dataset:
            domain_table = f"{bq_dataset}.{table_name}"
        elif source_type == "postgresql":
            domain_table = f"{pg_schema}.{table_name}"
        else:
            domain_table = table_name

        row_count = row_counts.get(table_name) or max((len(values) for values in columns.values()), default=0)
        for column, values in columns.items():
            existing = (await db.execute(
                select(MetadataRecord).where(
                    MetadataRecord.project_id == project_id,
                    MetadataRecord.data_domain_table == domain_table,
                    MetadataRecord.data_attribute == column,
                )
            )).scalar_one_or_none()

            sensitivity = _classify_sensitivity(column)
            business_term = _expand_business_term(column)
            sample = _get_sample_data(values)
            data_type = _detect_data_type(values)
            primary_key = _is_primary_key_candidate(values)
            nullable = _is_nullable(values)
            standard_format = _assess_standard_format(values)
            distinct_vals = _get_distinct_values(values, standard_format)

            if existing:
                existing.data_sensitivity = sensitivity
                existing.business_term = business_term
                existing.sample_data = sample
                existing.data_type = data_type
                existing.is_primary_key = primary_key
                existing.is_nullable = nullable
                existing.source_row_count = row_count
                existing.standard_format = standard_format
                existing.distinct_values = distinct_vals
                existing.updated_date = today
                existing.updated_by = initiated_by
                updated += 1
            else:
                next_seq += 1
                db.add(MetadataRecord(
                    project_id=project_id,
                    seq_no=next_seq,
                    business_users=customer_name or "",
                    data_domain_table=domain_table,
                    line_of_business=line_of_business,
                    table_type="Source",
                    project_name=project_name or "",
                    project_year=project_year or current_year,
                    data_steward=steward_str,
                    data_owner=owner_str,
                    data_attribute=column,
                    data_year=current_year,
                    data_sensitivity=sensitivity,
                    business_term=business_term,
                    definition_status="pending",
                    sample_data=sample,
                    data_type=data_type,
                    is_primary_key=primary_key,
                    is_nullable=nullable,
                    data_level="Raw",
                    standard_format=standard_format,
                    distinct_values=distinct_vals,
                    remarks="-",
                    source_type=source_type,
                    source_row_count=row_count,
                    updated_date=today,
                    updated_by=initiated_by,
                ))
                created += 1

    return {
        "project_id": str(project_id),
        "created": created,
        "updated": updated,
        "processed": created + updated,
        "tables": list(tables_data.keys()),
    }
